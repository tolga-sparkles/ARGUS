import io
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

TRT_OFFSET = timedelta(hours=3)


def _trt_now():
    return datetime.now(timezone.utc) + TRT_OFFSET


def trt_text():
    return _trt_now().strftime("%d.%m.%Y %H:%M:%S TRT")


COLUMNS = {
    "hostname": {"label": "Hostname"},
    "display_name": {"label": "Display Name"},
    "ip": {"label": "IP Address"},
    "status": {"label": "Status"},
    "os": {"label": "Operating System"},
    "site": {"label": "Site"},
    "role": {"label": "Role"},
    "environment": {"label": "Environment"},
    "cpu_percent": {"label": "CPU %"},
    "cpu_cores": {"label": "CPU Cores"},
    "memory_percent": {"label": "Memory %"},
    "memory_used_text": {"label": "Memory Used"},
    "memory_total_text": {"label": "Memory Total"},
    "disk_usage_percent": {"label": "Disk Used %"},
    "disk_free_percent": {"label": "Disk Free %"},
    "disk_used_text": {"label": "Disk Used"},
    "disk_total_text": {"label": "Disk Total"},
    "ping_latency_ms": {"label": "Ping (ms)"},
    "uptime_text": {"label": "Uptime"},
    "last_seen_text": {"label": "Last Seen"},
    "telemetry_state": {"label": "Telemetry State"},
    "datacenter": {"label": "Datacenter"},
    "cluster": {"label": "Cluster"},
    "virt_host": {"label": "Virtualization Host"},
    "node_type": {"label": "VM / Physical"},
}

PRESETS = {
    "standard": [
        "hostname", "ip", "status", "os", "role", "cpu_percent", "cpu_cores",
        "memory_percent", "memory_used_text", "memory_total_text",
        "disk_usage_percent", "disk_used_text", "disk_total_text",
        "ping_latency_ms", "last_seen_text",
    ],
    "inventory": [
        "hostname", "display_name", "ip", "os", "site", "role", "environment",
        "datacenter", "cluster", "virt_host",
    ],
    "resource_usage": [
        "hostname", "status", "cpu_percent", "cpu_cores", "memory_percent",
        "memory_used_text", "memory_total_text", "disk_usage_percent",
        "disk_used_text", "disk_total_text", "ping_latency_ms", "uptime_text",
    ],
    "technical": list(COLUMNS.keys()),
}

PDF_COLUMNS = ["hostname", "ip", "status", "cpu_percent", "memory_percent", "disk_usage_percent", "ping_latency_ms", "role"]

STATUS_LABELS = {
    "healthy": "HEALTHY",
    "warning": "WARNING",
    "critical": "CRITICAL",
    "offline": "OFFLINE",
    "no_telemetry": "NO TELEMETRY",
    "unknown": "UNKNOWN",
}

FILTER_LABELS = {
    "status": "Status",
    "site": "Site",
    "role": "Role",
    "os": "Operating System",
    "environment": "Environment",
    "q": "Search",
    "datacenter": "Datacenter",
    "cluster": "Cluster",
    "host": "Virtualization Host",
    "type": "Node Type",
}


def _value_for(item, key, virt_info):
    if key == "status":
        return STATUS_LABELS.get(item.get("state"), str(item.get("state", "")).upper())
    if key == "disk_free_percent":
        used = item.get("disk_usage_percent")
        if used is None:
            return None
        return round(100.0 - used, 1)
    if key == "node_type":
        if virt_info and virt_info.get("kind") == "vm":
            return "Virtual"
        if virt_info and virt_info.get("kind") == "host":
            return "Host"
        return "Physical"
    if key == "datacenter":
        return virt_info.get("datacenter") if virt_info else None
    if key == "cluster":
        return virt_info.get("cluster") if virt_info else None
    if key == "virt_host":
        return virt_info.get("host") if virt_info else None
    if key in ("cpu_percent", "memory_percent", "disk_usage_percent", "ping_latency_ms", "uptime_text"):
        value = item.get(key)
        return value if value is not None else None
    return item.get(key)


def _format_cell(value):
    if value is None or value == "":
        return "—"
    if isinstance(value, float):
        return int(value) if value.is_integer() else value
    return str(value)


def resolve_columns(report_type, custom_columns):
    if custom_columns:
        keys = [key for key in custom_columns if key in COLUMNS]
        if keys:
            return keys
    if report_type == "custom":
        return PRESETS["standard"]
    if report_type in PRESETS:
        return PRESETS[report_type]
    return PRESETS["standard"]


def resolve_dataset(items, virt_lookup, filters):
    rows = []
    for item in items:
        virt_info = None
        if virt_lookup:
            virt_info = virt_lookup.get(item["hostname"].lower()) or virt_lookup.get(item["ip"])
        rows.append((item, virt_info))

    applied = []
    for key, label in FILTER_LABELS.items():
        value = (filters or {}).get(key)
        if value and value not in ("", "all"):
            applied.append((label, str(value)))

    summary = {"total": len(rows), "healthy": 0, "warning": 0, "critical": 0, "offline": 0, "no_telemetry": 0}
    for item, _ in rows:
        state = item.get("state")
        if state in summary:
            summary[state] += 1
    return rows, applied, summary


def build_excel(rows, columns, meta):
    wb = Workbook()
    ws = wb.active
    ws.title = "Servers"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="161B22", end_color="161B22", fill_type="solid")
    status_fills = {
        "HEALTHY": PatternFill(start_color="1B3A2A", end_color="1B3A2A", fill_type="solid"),
        "WARNING": PatternFill(start_color="3A2F12", end_color="3A2F12", fill_type="solid"),
        "CRITICAL": PatternFill(start_color="3A1B1B", end_color="3A1B1B", fill_type="solid"),
        "OFFLINE": PatternFill(start_color="3A1B1B", end_color="3A1B1B", fill_type="solid"),
        "NO TELEMETRY": PatternFill(start_color="26282E", end_color="26282E", fill_type="solid"),
    }

    labels = [COLUMNS[key]["label"] for key in columns]
    ws.append(labels)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")

    for item, virt_info in rows:
        values = [_format_cell(_value_for(item, key, virt_info)) for key in columns]
        ws.append(values)
        status_idx = columns.index("status") + 1 if "status" in columns else None
        if status_idx:
            fill = status_fills.get(str(values[status_idx - 1]))
            if fill:
                ws.cell(row=ws.max_row, column=status_idx).fill = fill

    for index, key in enumerate(columns, start=1):
        width = max(len(COLUMNS[key]["label"]) + 2, 14)
        if key in ("hostname", "display_name", "ip", "datacenter", "cluster", "virt_host"):
            width = max(width, 22)
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet("Report Info")
    info_rows = [
        ("Generated At", meta["generated_at"]),
        ("Generated By", meta["generated_by"]),
        ("Total Servers", meta["total"]),
        ("Metric Scope", "Latest available telemetry"),
        ("", ""),
        ("Filters", ""),
    ]
    if meta["filters"]:
        for label, value in meta["filters"]:
            info_rows.append((label, value))
    else:
        info_rows.append(("Scope", "All servers"))
    for row in info_rows:
        info.append(row)
    info["A1"].font = Font(bold=True)
    info["A6"].font = Font(bold=True)
    for index, row in enumerate(info_rows, start=1):
        if index >= 7 and row[0]:
            info.cell(row=index, column=1).font = Font(bold=True)
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 46

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_csv(rows, columns):
    labels = [COLUMNS[key]["label"] for key in columns]
    lines = [",".join('"{0}"'.format(label.replace('"', '""')) for label in labels)]
    for item, virt_info in rows:
        values = [str(_format_cell(_value_for(item, key, virt_info))).replace('"', '""') for key in columns]
        lines.append(",".join('"{0}"'.format(value) for value in values))
    return ("\r\n".join(lines) + "\r\n").encode("utf-8-sig")


def build_pdf(rows, summary, meta):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=28, rightMargin=28, topMargin=26, bottomMargin=26,
        title="Server Status Report",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleCustom", parent=styles["Title"], fontSize=15, textColor=colors.HexColor("#f0f6fc"), spaceAfter=1)
    brand_style = ParagraphStyle("BrandCustom", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#8b949e"), spaceAfter=2)
    small = ParagraphStyle("SmallCustom", parent=styles["Normal"], fontSize=8.5, textColor=colors.HexColor("#8b949e"))
    section = ParagraphStyle("SectionCustom", parent=styles["Normal"], fontSize=10, textColor=colors.HexColor("#58a6ff"), spaceBefore=8, spaceAfter=4)

    story = [
        Paragraph("EXAMPLE ORGANIZATION", brand_style),
        Paragraph("ARGUS · Server Status Report", title_style),
        Spacer(1, 4),
        Paragraph("Generated {0} · Scope {1} · Servers {2}".format(
            meta["generated_at"], meta["scope_text"], summary["total"]), small),
        Paragraph("SUMMARY", section),
    ]

    summary_data = [
        ["TOTAL", "HEALTHY", "WARNING", "CRITICAL", "OFFLINE", "NO TELEMETRY"],
        [str(summary["total"]), str(summary["healthy"]), str(summary["warning"]),
         str(summary["critical"]), str(summary["offline"]), str(summary["no_telemetry"])],
    ]
    summary_table = Table(summary_data, colWidths=[70] * 6)
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#21262d")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#8b949e")),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("FONTSIZE", (0, 1), (-1, 1), 13),
        ("TEXTCOLOR", (0, 1), (-1, 1), colors.HexColor("#f0f6fc")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(summary_table)
    story.append(Paragraph("SERVERS", section))

    labels = [COLUMNS[key]["label"] for key in PDF_COLUMNS]
    data = [labels]
    for item, virt_info in rows:
        data.append([_format_cell(_value_for(item, key, virt_info)) for key in PDF_COLUMNS])

    col_widths = [110, 88, 72, 44, 44, 44, 52, 60]
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table_style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#21262d")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#8b949e")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("TEXTCOLOR", (0, 1), (-1, -1), colors.HexColor("#e1e4e8")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#21262d")),
        ("TOPPADDING", (0, 0), (-1, -1), 3.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
    ]
    for row_idx in range(1, len(data)):
        status = data[row_idx][2]
        if status == "CRITICAL":
            table_style.append(("TEXTCOLOR", (0, row_idx), (-1, row_idx), colors.HexColor("#f85149")))
        elif status == "WARNING":
            table_style.append(("TEXTCOLOR", (0, row_idx), (-1, row_idx), colors.HexColor("#d29922")))
        elif status == "OFFLINE":
            table_style.append(("BACKGROUND", (0, row_idx), (-1, row_idx), colors.HexColor("#2a1516")))
    table.setStyle(TableStyle(table_style))
    story.append(table)

    def footer(canvas, doc_obj):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#6e7681"))
        canvas.drawCentredString(landscape(A4)[0] / 2, 12, "Page {0} · ARGUS".format(doc_obj.page))
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def generate_filename(fmt, filters):
    now = _trt_now()
    parts = ["infrastructure"]
    for key, value in (filters or {}).items():
        if key in ("os", "status", "site", "role") and value and value not in ("", "all"):
            parts.append(str(value).lower().replace(" ", "_")[:16])
    parts.append(now.strftime("%Y-%m-%d"))
    return "{0}.{1}".format("_".join(parts), fmt)


def scope_text(scope, filters):
    if scope == "all":
        return "All servers"
    if scope == "selected":
        return "Selected servers"
    if filters:
        return ", ".join("{0} {1}".format(label, value) for label, value in _describe_filters(filters))
    return "All servers"


def _describe_filters(filters):
    out = []
    for key, label in FILTER_LABELS.items():
        value = (filters or {}).get(key)
        if value and value not in ("", "all"):
            out.append((label, str(value)))
    return out
